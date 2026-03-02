import logging
from pathlib import Path
from xml.sax.saxutils import escape

from jinja2 import Environment, FileSystemLoader

logger = logging.getLogger(__name__)


class DocumentService:
    def __init__(self) -> None:
        # 템플릿 디렉토리 설정 (app/templates)
        template_dir = Path(__file__).resolve().parent.parent / "templates"
        self.env = Environment(
            loader=FileSystemLoader(str(template_dir)),
            autoescape=True,
        )

    def _register_korean_font(self) -> str:
        """한글 폰트를 등록하고 폰트명을 반환합니다 (크로스 플랫폼)."""
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        font_name = "Helvetica"
        try:
            font_paths = [
                Path("C:/Windows/Fonts/malgun.ttf"),  # Windows
                Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"),  # Linux
                Path("/System/Library/Fonts/AppleSDGothicNeo.ttc"),  # Mac
            ]
            for path in font_paths:
                if path.exists():
                    if path.suffix.lower() == ".ttc":
                        pdfmetrics.registerFont(
                            TTFont("KoreanFont", str(path), subfontIndex=0)
                        )
                    else:
                        pdfmetrics.registerFont(TTFont("KoreanFont", str(path)))
                    font_name = "KoreanFont"
                    break
        except Exception as e:
            logger.warning("폰트 로드 실패: %s", e)

        return font_name

    def generate_demand_letter(
        self,
        data: dict[str, object],
        format: str = "text",
        output_path: str | None = None
    ) -> str:
        """내용증명 생성"""
        # 필수 필드 검증
        required_fields = ["recipient", "sender", "content"]
        missing = [field for field in required_fields if field not in data]
        if missing:
            raise ValueError(f"필수 데이터 누락: {', '.join(missing)}")

        if format == "text":
            template = self.env.get_template("demand_letter.j2")
            return template.render(**data)

        elif format == "pdf":
            if not output_path:
                raise ValueError("PDF 생성 시 output_path는 필수입니다.")

            from reportlab.lib.pagesizes import A4
            from reportlab.pdfgen import canvas

            font_name = self._register_korean_font()

            c = canvas.Canvas(output_path, pagesize=A4)
            c.setFont(font_name, 12)

            # 텍스트 렌더링 (줄바꿈 처리 필요)
            text = self.env.get_template("demand_letter.j2").render(**data)
            y = 800
            for line in text.split('\n'):
                c.drawString(50, y, line)
                y -= 20

            c.save()
            return output_path

        else:
            raise ValueError(f"지원하지 않는 포맷입니다: {format}")

    def generate_pdf_from_text(self, text: str, output_path: str) -> str:
        """텍스트 내용을 PDF로 생성 (Platypus 기반, CJK 자동 줄바꿈)"""
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer

        font_name = self._register_korean_font()

        body_style = ParagraphStyle(
            name="Korean",
            fontName=font_name,
            fontSize=11,
            leading=18,
            wordWrap="CJK",
            spaceAfter=4,
        )
        title_style = ParagraphStyle(
            name="KoreanTitle",
            fontName=font_name,
            fontSize=14,
            leading=22,
            wordWrap="CJK",
            spaceAfter=8,
            alignment=1,  # center
        )

        doc = SimpleDocTemplate(
            output_path,
            pagesize=A4,
            leftMargin=25 * mm,
            rightMargin=25 * mm,
            topMargin=25 * mm,
            bottomMargin=25 * mm,
        )

        story: list[Flowable] = []
        for line in text.split("\n"):
            if not line.strip():
                story.append(Spacer(1, 6))
                continue

            clean_line = line.strip()
            # ** 마크다운 볼드 → 제목 스타일로 처리
            if clean_line.startswith("**") and clean_line.endswith("**") and len(clean_line) > 4:
                clean_line = clean_line[2:-2]
                safe_line = escape(clean_line)
                story.append(Paragraph(safe_line, title_style))
            else:
                safe_line = escape(clean_line)
                story.append(Paragraph(safe_line, body_style))

        doc.build(story)
        return output_path

    def generate_hwpx_from_text(self, text: str, output_path: str) -> str:
        """텍스트 내용을 HWPX(한글) 파일로 생성"""
        from hwpx.document import HwpxDocument

        # 빈 문서 생성 (HwpxDocument.new() 사용)
        doc = HwpxDocument.new()

        # 텍스트를 문단별로 추가
        for line in text.split('\n'):
            doc.add_paragraph(line)

        # 파일 저장
        doc.save(output_path)
        return output_path

    def generate_docx_from_text(self, text: str, output_path: str) -> str:
        """텍스트 내용을 DOCX(워드) 파일로 생성 - 한글에서도 열 수 있음"""
        from docx import Document
        from docx.shared import Pt

        doc = Document()

        # 텍스트를 문단별로 추가
        for line in text.split('\n'):
            para = doc.add_paragraph(line)
            # 기본 폰트 설정
            for run in para.runs:
                run.font.name = '맑은 고딕'
                run.font.size = Pt(11)

        # 파일 저장
        doc.save(output_path)
        return output_path
